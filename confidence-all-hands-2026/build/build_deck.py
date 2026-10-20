"""Build the Confidence 2026 Year-End All Hands overview on the VML 2026 template.

Usage: python3 build_deck.py [output.pptx]

Every slide sits on a native VML layout (light and dark masters, header rules, spark mark,
slide number), and all type uses the template's Snowflake Sans theme fonts. Color follows the
brand rule: black and white first, with one supporting hue (Gold) used sparingly.
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "template" / "VML-Template-Feb-2026.pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "Confidence-2026-All-Hands-Overview.pptx"

# --- VML palette (template "VML" scheme + brand grayscale + one secondary hue: Gold) ----
BLACK = RGBColor(0x19, 0x19, 0x19)
WHITE = RGBColor(0xF5, 0xF5, 0xF5)
DGRAY = RGBColor(0x36, 0x37, 0x32)
MGRAY = RGBColor(0x53, 0x54, 0x4A)
GRAY = RGBColor(0x8B, 0x8C, 0x81)
LGRAY = RGBColor(0xC2, 0xC4, 0xB8)
GOLD = RGBColor(0xD3, 0xA3, 0x4A)
LGOLD = RGBColor(0xF9, 0xD2, 0x8C)
DGOLD = RGBColor(0xB3, 0x7A, 0x00)

BODY = "+mn-lt"   # Snowflake Sans Book (theme minor font)
HEAD = "+mj-lt"   # Snowflake Sans Medium (theme major font)

# Content grid (inches), matching the template's header rules at 0.55" and 1.61"
L, R = 0.29, 13.05
W = R - L
TOP = 1.95

prs = Presentation(str(TEMPLATE))
LIGHT, DARK = prs.slide_masters[0], prs.slide_masters[1]


def layout(master, name):
    for lay in master.slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError(name)


# --- remove the template's guidance and sample slides ----------------------------------
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


# --- helpers ---------------------------------------------------------------------------
def _inherit_slide_number(slide, lay):
    """python-pptx skips slide-number placeholders; clone it from the layout."""
    for ph in lay.placeholders:
        if ph.placeholder_format.type == PH.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))


def new_slide(master, name, title=None, notes=None, keep=()):
    lay = layout(master, name)
    s = prs.slides.add_slide(lay)
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if ph.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE) and title is not None:
            ph.text_frame.text = title
        elif idx in keep:
            continue
        else:
            ph._element.getparent().remove(ph._element)
    _inherit_slide_number(s, lay)
    if notes:
        s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def content(title, notes, dark=False):
    """Interior 1.0: small header between the two template rules, open content area below."""
    return new_slide(DARK if dark else LIGHT, f"Interior 1.0 ({'Dark' if dark else 'Light'})",
                     title=title, notes=notes)


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def _style_run(run, size, color, font, bold=False, italic=False):
    f = run.font
    f.size = Pt(size)
    f.name = font
    f.bold = bold
    f.italic = italic
    if color is not None:
        f.color.rgb = color


def text(slide, x, y, w, h, paras, size=14, color=BLACK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=0, line=None, name=None):
    """paras: str | list of str | list of dict(t=..., size, color, font, bold, italic, align, after, check)
    A dict's t may itself be a list of run dicts for mixed styling. check=True adds the VML □ bullet."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        if not isinstance(p, dict):
            p = {"t": p}
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("after", after))
        if p.get("before"):
            para.space_before = Pt(p["before"])
        if line or p.get("line"):
            para.line_spacing = p.get("line", line)
        runs = p["t"] if isinstance(p["t"], list) else [{"t": p["t"]}]
        for rd in runs:
            r = para.add_run()
            r.text = rd["t"]
            _style_run(r, rd.get("size", p.get("size", size)), rd.get("color", p.get("color", color)),
                       rd.get("font", p.get("font", font)), rd.get("bold", p.get("bold", bold)),
                       rd.get("italic", p.get("italic", False)))
        if p.get("check"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.3))))
            pPr.set("indent", str(-int(Inches(0.3))))
            etree.SubElement(pPr, qn("a:buFont"), typeface=BODY)
            etree.SubElement(pPr, qn("a:buChar"), char="□")
    return tb


def box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def rule(slide, x1, y1, x2, y2, color=LGRAY, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def label_in(shape, t, size=12, color=WHITE, font=HEAD, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    _style_run(r, size, color, font)


def table(slide, x, y, w, col_w, rows, row_h=0.5, size=12, head_color=BLACK, body_color=BLACK,
          rule_color=LGRAY, highlight=None):
    """Editable native table in the VML style: no fills, hairline row rules, bold header row."""
    gf = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            if highlight is not None and ri == highlight:
                c.fill.solid()
                c.fill.fore_color.rgb = LGOLD
            else:
                c.fill.background()
            c.margin_left = Inches(0.08 if ci else 0)
            c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = val
            head = ri == 0
            _style_run(r, size, head_color if head else body_color, HEAD if (head or ci == 0) else BODY)
            tcPr = c._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side), w=str(int(Pt(1 if (head and side == "a:lnB") else 0.75))))
                if side == "a:lnB":
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=str(head_color if head else rule_color))
                else:
                    etree.SubElement(ln, qn("a:noFill"))
            # schema order: borders must precede the cell fill
            for fill_tag in ("a:solidFill", "a:noFill"):
                for el in tcPr.findall(qn(fill_tag)):
                    tcPr.remove(el)
                    tcPr.append(el)
    return gf


def numeral(slide, x, y, n, color=BLACK, size=40):
    return text(slide, x, y, 2.2, 0.7 if size < 60 else 1.2, f"{n:02d}", size=size, color=color, font=HEAD)






# =======================================================================================
# 1 — OVERVIEW
# =======================================================================================
s = content("Year-End All Hands  |  Overview for feedback", """
Our final Confidence All Hands of 2026 takes a slightly different shape this year: a 2026 Impact Report and a 2027 Future Outlook, rather than a traditional year-end celebration.

Why: there's a lot worth recognizing, but this has been a demanding year. The meeting should show what we accomplished, appreciate the people behind it, and look ahead with some curiosity, without feeling overly celebratory.

It's built around three parts: Our Impact, Our People and Our Future. Confidence leadership updates keep a dedicated slot.

We're sharing this to get your reactions before planning goes further.
""", dark=True)
text(s, L, TOP, W, 0.8, "Confidence 2026: Our Impact, Our Future", size=44, color=WHITE, font=HEAD)
text(s, L, TOP + 0.95, 11.0, 0.9, "Our year-end All Hands, shaped as a 2026 Impact Report and a 2027 Future Outlook "
     "rather than a traditional celebration.", size=20, color=LGRAY, line=1.1)
pillars = [
    ("Our Impact", "What Confidence delivered for our clients and our practice this year"),
    ("Our People", "The individuals and teams behind that work"),
    ("Our Future", "What our practice could look like in 2027 and beyond"),
]
cw = (W - 2 * 0.4) / 3
for i, (h, d) in enumerate(pillars):
    x = L + i * (cw + 0.4)
    rule(s, x, 4.0, x + cw, 4.0, color=MGRAY)
    text(s, x, 4.2, 0.8, 0.5, f"{i + 1:02d}", size=16, color=GOLD, font=HEAD)
    text(s, x, 4.65, cw - 0.1, 0.5, h, size=24, color=WHITE, font=HEAD)
    text(s, x, 5.25, cw - 0.1, 0.9, d, size=15, color=LGRAY, line=1.15)
text(s, L, 6.55, W, 0.4, [{"t": [
    {"t": "60 minutes  ·  late November, after November reporting  ·  ", "color": LGRAY},
    {"t": "includes Confidence leadership updates", "color": LGRAY},
]}], size=14)

# =======================================================================================
# 2 — FORMAT
# =======================================================================================
s = content("Format: a proposed 60-minute flow", """
A proposed run of show. Timing is a starting point.

The flow moves from what we did, to who made it happen, to what could come next. Leadership updates keep a dedicated 10 minutes.

The popcorn conversation is the most different part of the format: a short, live discussion where a few people from across Confidence build on each other's ideas about the future.
""")
segs = [
    (5, "Welcome", "Acknowledge the year and set the tone", "frame"),
    (10, "Leadership updates", "Two or three concise updates on the business, clients and practice", "frame"),
    (15, "Our Impact", "Highlights of what Confidence delivered in 2026", "core"),
    (8, "Our People", "Short spotlights on individuals and teams", "core"),
    (5, "Our Future", "Three provocations about where our work is heading", "core"),
    (12, "Popcorn conversation", "A live, “yes, and” discussion that builds on those ideas", "key"),
    (5, "Close", "Appreciation and a few takeaways", "frame"),
]
fills = {"frame": (LGRAY, BLACK), "core": (BLACK, WHITE), "key": (GOLD, BLACK)}
unit = W / 60
x = L
bar_y, bar_h = TOP + 0.05, 0.55
for mins, name, desc, kind in segs:
    f, c = fills[kind]
    seg = box(s, x, bar_y, mins * unit - 0.04, bar_h, fill=f)
    label_in(seg, f"{mins} min", size=12, color=c)
    x += mins * unit
ty, rows_h = 3.0, 0.57
t0 = 0
for i, (mins, name, desc, kind) in enumerate(segs):
    y = ty + i * rows_h
    f, _ = fills[kind]
    box(s, L, y + 0.18, 0.2, 0.2, fill=f, line=BLACK if kind == "frame" else None)
    text(s, L + 0.45, y + 0.13, 1.3, 0.35, f"{t0:02d}–{t0 + mins:02d}", size=14, color=MGRAY)
    text(s, L + 1.8, y + 0.13, 3.6, 0.35, name, size=15, font=HEAD)
    text(s, L + 5.5, y + 0.13, W - 5.5, 0.35, desc, size=15)
    rule(s, L, y + rows_h, R, y + rows_h)
    t0 += mins

# =======================================================================================
# 3 — CONTENT
# =======================================================================================
s = content("Content: what each part could include", """
A quick look at what each part could include. These are examples and thought starters, not final selections.

Our Impact: a few strong metrics and short client stories, organized by theme rather than by client, so it reads as one practice story rather than a status update. Possible featured stories from this year's Technology Reports include United Rentals (AI-enabled delivery and automation), Sherwin-Williams (shared delivery standards across the T4 product teams) and Mazda (program-level quality governance). A visual summary can show the breadth of work across other accounts.

Our People: short spotlights connected to real contributions, such as innovation, leadership through tough delivery, collaboration, mentorship and knowledge sharing. Nominations would come from group and pod leads. Not a competitive awards format.

Our Future: three "what if" provocations inspired by VML Intelligence's Future 100: 2026, each tied to something we did this year. Then a 12-minute popcorn conversation: four or five people from different disciplines and pods react to a scenario and build on each other's ideas, with the audience joining in chat.
""")
parts = [
    ("Our Impact", "15 min", "A few strong metrics and short client stories that read as one practice story.", [
        "Featured stories, e.g. United Rentals, Sherwin-Williams, Mazda",
        "Organized by theme: AI in delivery, quality at scale, shared practices",
        "A visual summary of highlights across other accounts",
    ]),
    ("Our People", "8 min", "Recognition connected to real contributions, not an awards show.", [
        "Short spotlights on individuals and teams",
        "Innovation, leadership, collaboration, mentorship",
        "Nominations from group and pod leads",
    ]),
    ("Our Future", "17 min", "Provocations inspired by VML's Future 100, then a live conversation.", [
        "What if AI becomes a member of every delivery team?",
        "What if trust becomes our most valuable deliverable?",
        "Popcorn conversation: 4–5 voices across disciplines, “yes, and”",
    ]),
]
cw = (W - 2 * 0.3) / 3
ch = 7.0 - TOP
for i, (name, mins, what, items) in enumerate(parts):
    x = L + i * (cw + 0.3)
    box(s, x, TOP, cw, ch, line=BLACK)
    text(s, x + 0.25, TOP + 0.2, 1.0, 0.4, f"{i + 1:02d}", size=16, color=DGOLD if i == 2 else BLACK, font=HEAD)
    text(s, x + cw - 1.45, TOP + 0.22, 1.2, 0.4, mins, size=14, color=MGRAY, align=PP_ALIGN.RIGHT)
    text(s, x + 0.25, TOP + 0.7, cw - 0.5, 0.45, name, size=24, font=HEAD)
    text(s, x + 0.25, TOP + 1.3, cw - 0.5, 0.9, what, size=15, line=1.1)
    rule(s, x + 0.25, TOP + 2.4, x + cw - 0.25, TOP + 2.4)
    text(s, x + 0.25, TOP + 2.55, cw - 0.5, 0.3, "EXAMPLES", size=10, color=MGRAY, font=HEAD)
    text(s, x + 0.25, TOP + 2.9, cw - 0.5, 2.1, [{"t": it, "check": True} for it in items],
         size=14, after=8, line=1.05)

# =======================================================================================
# 4 — FEEDBACK
# =======================================================================================
s = content("We'd like your reactions", """
Open it up. The goal is honest reactions, not sign-off.

Prompts if the conversation needs a nudge:
• Does an impact-and-future format feel right for the year we've had?
• Is the time balanced well across updates, impact, recognition and discussion?
• Are there stories, accomplishments or people we should make sure to include?
• Does a live popcorn conversation sound engaging, or would something else work better? Who would be great in it?
""", dark=True)
qs = [
    ("The format", "Does an impact-and-future All Hands feel right for the year we've had?"),
    ("The balance", "Is time spread well across updates, impact, recognition and discussion?"),
    ("The content", "Which stories, accomplishments or people should we be sure to include?"),
    ("The conversation", "Would a live popcorn conversation land with the team? Who would be great in it?"),
]
cw, rh = (W - 0.5) / 2, 2.2
for i, (h, q) in enumerate(qs):
    col, row = i % 2, i // 2
    x = L + col * (cw + 0.5)
    y = TOP + 0.1 + row * rh
    text(s, x, y, 0.9, 0.6, f"{i + 1:02d}", size=28, color=GOLD, font=HEAD)
    text(s, x + 1.0, y + 0.05, cw - 1.0, 0.45, h, size=22, color=WHITE, font=HEAD)
    text(s, x + 1.0, y + 0.6, cw - 1.0, 0.9, q, size=16, color=LGRAY, line=1.15)
    rule(s, x, y + rh - 0.3, x + cw, y + rh - 0.3, color=MGRAY)

prs.core_properties.title = "Confidence 2026: Our Impact, Our Future"
prs.core_properties.subject = "Year-End All Hands overview for feedback"
prs.core_properties.author = "Zack Jemas"
prs.save(str(OUT))
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
